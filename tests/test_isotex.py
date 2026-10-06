"""Isotex blocks: an alternative way to build, chosen in the brief; the assistant counts the pieces."""
import os

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')

import pytest

from archforge.architecture.wall_types import WALL_TYPES, indicative_u, total_thickness
from archforge.core.commands import AddEntity, CommandStack, UpdateEntity
from archforge.core.model import Document, Entity
from archforge.construction.isotex import (BLOCKS, BLOCKS_PER_M2, concrete_per_m2, count_wall, take_off_isotex,
                                           wall_type_of)


def _wall(stack, x1, y1, x2, y2, code='HB 30/19', z=0.0, h=3.0):
    t = wall_type_of(code)
    w = Entity('wall', {'x1': x1, 'y1': y1, 'x2': x2, 'y2': y2, 'z': z, 'height': h,
                        'thickness': total_thickness(t), 'wall_type': t}, name='Τ')
    stack.execute(AddEntity(w))
    return w


def test_catalogue_with_and_without_insulation_as_wall_types():
    assert BLOCKS_PER_M2 == pytest.approx(8.0)
    assert any(v[2] == 0 for v in BLOCKS.values()) and any(v[2] > 0 for v in BLOCKS.values())
    for code, (_l, total, eps, core, *_r) in BLOCKS.items():
        t = wall_type_of(code)
        assert t in WALL_TYPES and total_thickness(t) == pytest.approx(total)
    assert indicative_u(wall_type_of('HB 25/16')) == pytest.approx(0.79)          # published U wins
    assert concrete_per_m2('HB 30/19') == (151.0, False)
    litres, est = concrete_per_m2('HDIII 38/14')
    assert est and litres == pytest.approx(150 * 0.79)


def test_pieces_per_wall_net_of_openings_with_waste_and_concrete():
    doc = Document(); st = CommandStack(doc)
    w = _wall(st, 0, 0, 5, 0)                                   # 5 × 3 = 15 m²
    st.execute(AddEntity(Entity('window', {'offset': 2.5, 'width': 1.0, 'height': 1.5, 'sill': 0.9}, parent_id=w.id)))
    r = count_wall(doc, w)
    assert r['net_m2'] == pytest.approx(13.5) and r['courses'] == 12
    assert r['blocks'] == 112                                   # ⌈13.5 × 8 × 1.03⌉ = ⌈111.24⌉
    assert r['jamb_cuts'] == 12                                 # 2 × 6 courses of the 1.5 m window
    assert r['concrete_m3'] == pytest.approx(13.5 * 0.151, abs=0.01)


def test_take_off_groups_by_block_and_storey_and_ignores_other_walls():
    doc = Document(); st = CommandStack(doc)
    _wall(st, 0, 0, 4, 0); _wall(st, 4, 0, 4, 4, code='HDIII 30/7'); _wall(st, 0, 0, 4, 0, z=3.0)
    st.execute(AddEntity(Entity('wall', {'x1': 0, 'y1': 4, 'x2': 4, 'y2': 4, 'z': 0, 'height': 3, 'thickness': 0.2})))
    r = take_off_isotex(doc)
    assert r['totals']['walls'] == 3 and set(r['by_code']) == {'HB 30/19', 'HDIII 30/7'}
    assert set(r['by_storey']) == {0.0, 3.0}
    assert r['totals']['blocks'] == sum(w['blocks'] for w in r['walls'])
    assert 'έλεγχο' in r['provenance']


def test_brief_offers_isotex_and_the_assistant_counts_pieces():
    from archforge.assistant.suggestions import propose
    from archforge.project.brief import INTERIOR_WALLS, WALL_SYSTEMS, load_bearing_walls, wall_defaults
    doc = Document(); st = CommandStack(doc)
    st.execute(AddEntity(Entity('project_brief', {'wall_system': wall_type_of('HDIII 38/14'),
                                                  'interior_walls': 'isotex_hb_25_16'})))
    assert wall_defaults(doc) == (wall_type_of('HDIII 38/14'), pytest.approx(0.38))
    assert load_bearing_walls(doc) and INTERIOR_WALLS['isotex_hb_25_16'][2]
    assert sum(k.startswith('isotex_') for k in WALL_SYSTEMS) == len(BLOCKS)
    _wall(st, 0, 0, 6, 0, code='HDIII 38/14')
    (p,) = [p for p in propose(doc) if p.key == 'I-1']
    assert 'τεμάχια' in p.title and p.run == 'isotex'


def test_menu_shows_the_take_off():
    from PySide6.QtWidgets import QApplication
    from archforge.ui.main_window import MainWindow
    app = QApplication.instance() or QApplication([])
    window = MainWindow(); window._no_modal_dialogs = True
    try:
        _wall(window.stack, 0, 0, 5, 0)
        result = window._show_isotex()
        assert result['totals']['blocks'] > 0 and 'HB 30/19' in window._isotex_view.toPlainText()
    finally:
        window._mark_clean(); window.close(); app.processEvents()


def test_isotex_has_its_own_priced_list_and_excel_computes_the_totals(tmp_path):
    import zipfile
    from archforge.quantities.materials import all_lists, isotex_list, write_workbook
    doc = Document(); st = CommandStack(doc)
    _wall(st, 0, 0, 5, 0); _wall(st, 5, 0, 5, 4, code='HDIII 38/14')
    rows = isotex_list(doc)
    assert [r[1] for r in rows[:2]] == ['τεμ.', 'τεμ.'] and any(r[2] is None for r in rows)   # reinforcement: from the study
    assert all_lists(doc)[0][0] == 'Isotex'
    path = write_workbook(doc, str(tmp_path / 'isotex.xlsx'), only=('Isotex',))
    with zipfile.ZipFile(path) as z:
        assert z.testzip() is None
        sheet = z.read('xl/worksheets/sheet1.xml').decode()
        assert 'IF(OR(C4=&quot;&quot;,D4=&quot;&quot;),0,C4*D4)' in sheet or 'C4*D4' in sheet
        assert f'SUM(E4:E{3 + len(rows)})' in sheet
        assert 'HDIII 38/14' in sheet and 'Τιμή μονάδας' in sheet
        assert 'name="Isotex"' in z.read('xl/workbook.xml').decode()


def test_priced_workbook_opens_in_a_spreadsheet_engine(tmp_path):
    from archforge.quantities.xlsx import write_priced_workbook
    path = write_priced_workbook(str(tmp_path / 'x.xlsx'), [('Α', 'Τίτλος', [('Τούβλο', 'τεμ.', 10), ('Άμμος', 'm³', None)], ())])
    try:
        import openpyxl
    except ImportError:
        pytest.skip('openpyxl not installed')
    ws = openpyxl.load_workbook(path)['Α']
    assert ws['A4'].value == 'Τούβλο' and ws['C4'].value == 10 and ws['E6'].value == '=SUM(E4:E5)'
